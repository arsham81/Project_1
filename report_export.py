"""Generate matching Excel and PDF reports without opening the input workbooks for writing."""
import io
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment
from app import read_plan, compare

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'pdf.fonttype':42})
INK='#173249';TEAL='#008c87';COPPER='#bd8054';GRAY='#afbdcb'

def fmt(n,d=0):
 return f'{n:,.{d}f}'.replace(',',' ').replace('.',',')

def prepare(first,second):
 paths=[Path(first),Path(second)]
 if any(not p.is_file() for p in paths):raise ValueError('Ընտրեք երկու գոյություն ունեցող Excel ֆայլեր։')
 if paths[0].resolve()==paths[1].resolve():raise ValueError('Ընտրեք երկու տարբեր ֆայլեր։')
 return compare(*(read_plan(p.read_bytes(),p.name) for p in paths))

def clean(ax):
 ax.set_facecolor('white');ax.spines[['top','right','left']].set_visible(False)
 ax.spines['bottom'].set_color('#d8e1e8');ax.tick_params(length=0,labelcolor=INK)
 ax.grid(axis='x',alpha=.12);ax.set_axisbelow(True)

def figures(result):
 values={m['key']:m for m in result['comparison']}
 fig,axes=plt.subplots(2,2,figsize=(11.69,8.27));fig.patch.set_facecolor('#f1f5f8')
 fig.suptitle('ԶՄՄԿ · Պլանների համեմատություն',fontsize=22,color=INK,weight='bold',y=.97)
 fig.text(.07,.91,'Բազային և թարմացված տարեկան ցուցանիշներ',fontsize=12,color=INK)
 for ax,key in zip(axes.flat,['cu','mo','processed','ore']):
  m=values.get(key)
  if not m:ax.axis('off');continue
  clean(ax);ax.barh([1,0],[m['base'],m['updated']],height=.45,color=[GRAY,TEAL])
  ax.set_yticks([1,0],['Բազային','Թարմացված']);ax.set_title(m['title'],loc='left',fontsize=13,weight='bold',color=INK,pad=15)
  mx=max(m['base'],m['updated'],1);ax.set_xlim(0,mx*1.35)
  for y,v in [(1,m['base']),(0,m['updated'])]:ax.text(v+mx*.025,y,fmt(v),va='center',fontsize=10)
  ax.set_xticks([]);ax.set_xlabel(m['unit'],color=INK)
 fig.subplots_adjust(left=.15,right=.97,bottom=.13,top=.82,wspace=.38,hspace=.65)
 fig.text(.07,.045,'Թարմացված տարեկան արժեքը կարող է ներառել փաստացի ամիսներ և մնացած ամիսների պլան։',fontsize=9,color=INK)
 yield fig

 fig,axes=plt.subplots(3,1,figsize=(11.69,8.27));fig.patch.set_facecolor('#f1f5f8')
 fig.suptitle('Փոփոխության գործոններ · Փաստ և պլան',fontsize=21,color=INK,weight='bold',y=.97)
 ax=axes[0];clean(ax)
 if result['factors']:
  fs=result['factors'];vs=[x['value'] for x in fs];ys=[2,1,0]
  ax.barh(ys,vs,color=[TEAL if v>=0 else COPPER for v in vs],height=.5)
  ax.set_yticks(ys,[x['title'] for x in fs]);ax.set_title('Պղնձի փոփոխության գործոնները · տ',loc='left',fontsize=13,weight='bold',color=INK)
  scale=max(max(abs(v) for v in vs),1);ax.set_xlim(min(0,min(vs))-scale*.25,max(0,max(vs))+scale*.3)
  for y,v in zip(ys,vs):ax.text(v+scale*.025 if v>=0 else v-scale*.025,y,('+' if v>0 else '')+fmt(v),va='center',ha='left' if v>=0 else 'right',fontsize=10)
  ax.set_xticks([])
 else:ax.axis('off');ax.text(0,.5,'Գործոնային բաժանման համար բավարար տվյալներ չկան։')
 for ax,key in zip(axes[1:],['cu','mo']):
  s=result['updated']['split'].get(key);clean(ax)
  title='Պղինձ' if key=='cu' else 'Մոլիբդեն'
  ax.set_title(title+' · փաստ և պլան',loc='left',fontsize=13,weight='bold',color=INK)
  if s and s['actual']>=0 and s['plan']>=0 and s['actual']+s['plan']>0:
   a,p=s['actual'],s['plan'];total=a+p
   ax.barh([0],[a],color=TEAL,height=.45,label='Փաստ');ax.barh([0],[p],left=[a],color=COPPER,height=.45,label='Պլան / սպասում')
   ax.set_xlim(0,total);ax.set_yticks([]);ax.set_xticks([])
   ax.text(0,-.45,f'Փաստ՝ {fmt(a)} տ · {a/total:.1%} · {s["actual_months"]} ամիս',fontsize=10,color=INK)
   ax.text(total,-.45,f'Պլան / սպասում՝ {fmt(p)} տ · {p/total:.1%}',ha='right',fontsize=10,color=INK)
   ax.set_ylim(-.7,.5);ax.legend(loc='upper right',frameon=False,ncol=2,fontsize=9)
  else:ax.axis('off');ax.text(.02,.4,'Ամսական բաժանումը հասանելի չէ կամ չի համընկնում տարեկան գումարին։',fontsize=11)
 fig.subplots_adjust(left=.26,right=.94,bottom=.17,top=.85,hspace=.95)
 fig.text(.06,.08,'Գործոնների հերթականությունը՝ ծավալ → պարունակություն → կորզում։',fontsize=9,color=INK)
 fig.text(.06,.055,'Օգտագործվում է արտադրությունից հաշվարկված տարեկան կորզումը, որը կարող է տարբերվել ֆայլի միջինից։',fontsize=9,color=INK)
 yield fig

 # Paginate rather than shrinking the complete comparison into unreadable text.
 metrics=result['comparison']
 for start in range(0,len(metrics),10):
  fig,ax=plt.subplots(figsize=(11.69,8.27));fig.patch.set_facecolor('white');ax.axis('off')
  ax.set_title('Համադրելի ցուցանիշներ',loc='left',fontsize=21,weight='bold',color=INK,pad=25)
  data=[]
  for m in metrics[start:start+10]:
   d=4 if m['unit']=='%' else 0
   data.append([m['title']+'\n('+m['unit']+')',fmt(m['base'],d),fmt(m['updated'],d),fmt(m['delta'],d)+(' տ.կ.' if m['unit']=='%' else ''),'—' if m['percent'] is None else fmt(m['percent'],2)+'%'])
  table=ax.table(cellText=data,colLabels=['Ցուցանիշ','Բազային','Թարմացված','Տարբերություն','Փոփոխություն'],colWidths=[.34,.165,.165,.18,.15],cellLoc='left',loc='upper left',bbox=[0,.13,1,.77])
  table.auto_set_font_size(False);table.set_fontsize(9)
  for (r,c),cell in table.get_celld().items():
   cell.set_edgecolor('#e5edf2')
   if r==0:cell.set_facecolor(INK);cell.set_text_props(color='white',weight='bold')
   else:cell.set_facecolor('#f1f5f8' if r%2 else 'white')
  fig.subplots_adjust(left=.045,right=.955,top=.89,bottom=.12)
  fig.text(.06,.065,'Բացարձակ տարբերությունները տոկոսային ցուցանիշների համար ներկայացված են տոկոսային կետերով։',fontsize=9,color=INK)
  yield fig

 fig=plt.figure(figsize=(11.69,8.27));fig.patch.set_facecolor('white')
 lines=['Աղբյուրներ և մեթոդաբանություն','']
 for label,p in [('Բազային',result['base']),('Թարմացված',result['updated'])]:
  lines.extend([label+':',p['filename'],'Թերթ՝ '+p['sheet'],''])
 lines.extend(['Օգտագործվել են Excel-ում պահպանված հաշվարկված արժեքները։','Ծրագիրը չի վերահաշվարկում Excel բանաձևերը։','Եթե արժեքները հնացած են, նախ վերահաշվարկեք և պահպանեք Excel ֆայլը։','','Գործոնների բաժանումը կախված է փոխարինման հերթականությունից։','Բացակայող տվյալները չեն փոխարինվում զրոներով։','Մուտքային Excel ֆայլերը չեն փոփոխվում։','','Աղբյուր բջիջների ամբողջական ցանկը՝ Excel հաշվետվության «Աղբյուրներ» թերթում։'])
 fig.text(.07,.92,'\n'.join(lines),va='top',fontsize=12,color=INK,linespacing=1.8,wrap=True)
 yield fig

def export_reports(result,excel_path,pdf_path):
 excel_path,pdf_path=Path(excel_path),Path(pdf_path)
 wb=Workbook();ws=wb.active;ws.title='Համեմատություն'
 ws.append(['Ցուցանիշ','Միավոր','Բազային','Թարմացված','Տարբերություն','Փոփոխություն, %'])
 for m in result['comparison']:
  ws.append([m['title'],m['unit'],m['base'],m['updated'],m['delta'],m['percent']])
 ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 for row in list(ws)[1:]:
  for cell in row[2:]:cell.number_format='#,##0.0000;[Red]-#,##0.0000'
 src=wb.create_sheet('Աղբյուրներ');src.append(['Տարբերակ','Ֆայլ','Թերթ','Ցուցանիշ','Բջիջ','Արժեք'])
 for label,p in [('Բազային',result['base']),('Թարմացված',result['updated'])]:
  for m in p['metrics'].values():src.append([label,p['filename'],p['sheet'],m['title'],m['cell'],m['value']])
 detail=wb.create_sheet('Փաստ և պլան');detail.append(['Ցուցանիշ','Փաստ','Պլան / սպասում','Փաստացի ամիսներ'])
 for key,s in result['updated']['split'].items():
  detail.append([result['updated']['metrics'][key]['title'],s['actual'],s['plan'],s['actual_months']])
 factors=wb.create_sheet('Գործոններ');factors.append(['Պղնձի փոփոխության գործոն','Տոննա'])
 for f in result['factors'] or []:factors.append([f['title'],f['value']])
 notes=wb.create_sheet('Նշումներ');notes.append(['Մեթոդաբանություն'])
 for line in ['Թարմացված տարեկան արժեքը կարող է ներառել փաստ և պլան։','Տոկոսային ցուցանիշների բացարձակ տարբերությունները տոկոսային կետերով են։','Excel բանաձևերը չեն վերահաշվարկվում. օգտագործվում են պահպանված արժեքները։','Գործոնները՝ ծավալ → պարունակություն → կորզում հերթականությամբ։','Գործոնային բաժանման համար օգտագործվում է արտադրությունից հաշվարկված տարեկան կորզումը։','Բացակայող ցուցանիշները բաց են թողնվում։']:
  notes.append([line])
 graph=wb.create_sheet('Գրաֆիկներ');buffers=[]
 try:
  with PdfPages(pdf_path) as pdf:
   for i,fig in enumerate(figures(result)):
    try:
     pdf.savefig(fig)
     if i<2:
      buffer=io.BytesIO();fig.savefig(buffer,format='png',dpi=160);buffer.seek(0);buffers.append(buffer)
      img=Image(buffer);img.width=1122;img.height=794;graph.add_image(img,'A'+str(1+i*43))
    finally:plt.close(fig)
  for sheet in wb:
   if sheet==graph:continue
   for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173249')
   for column in sheet.columns:
    sheet.column_dimensions[column[0].column_letter].width=min(75,max(18,max(len(str(c.value or '')) for c in column)+2))
  wb.save(excel_path)
 finally:
  for buffer in buffers:buffer.close()
 return excel_path,pdf_path
