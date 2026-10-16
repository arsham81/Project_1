"""Two-file ZCMC plan comparison. Run: python app.py"""
import base64
import io
import json
import math
import os
import re
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import openpyxl

METRICS = {
 'ore': ('Հանքաքարի արդյունահանում', 'տ', ['добыча руды']),
 'waste': ('Մակաբացման ապարներ', 'տ', ['добыча вскрышных пород']),
 'processed': ('Հանքաքարի վերամշակում', 'չոր տ', ['итого переработка руды']),
 'cu_grade': ('Cu պարունակություն', '%', ['среднее содержание cu в руде']),
 'mo_grade': ('Mo պարունակություն', '%', ['среднее содержание mo в руде']),
 'cu_recovery': ('Cu կորզում', '%', ['извлечение меди (cu)']),
 'mo_recovery': ('Mo կորզում', '%', ['извлечение молибдена (mo)']),
 'cu': ('Պղինձ խտանյութում', 'տ', ['количество меди (cu) в медном концентрате']),
 'mo': ('Մոլիբդեն խտանյութում', 'տ', ['количество молибдена (mo) в молибденовом концентрате']),
 'cu_conc': ('Պղնձի խտանյութ', 'խոնավ տ', ['выпуск медного (cu) концентрата']),
 'mo_conc': ('Մոլիբդենի խտանյութ', 'խոնավ տ', ['выпуск молибденового (mo) концентрата']),
 'loading': ('Ընդհանուր բեռնում', 'տ', ['общая погрузка']),
 'transport': ('Ընդհանուր փոխադրում', 'տ·կմ', ['общая перевозка']),
}

def norm(v):
 return re.sub(r'\s+', ' ', str(v or '').lower().replace('ё','е')).strip()

def read_plan(data, filename):
 if not filename.lower().endswith('.xlsx'):
  raise ValueError('Ընտրեք .xlsx ֆայլ։')
 if len(data)>20*1024*1024: raise ValueError('Ֆայլը գերազանցում է 20 ՄԲ սահմանը։')
 try:
  with zipfile.ZipFile(io.BytesIO(data)) as z:
   if sum(x.file_size for x in z.infolist())>100*1024*1024:
    raise ValueError('Excel ֆայլի բացված չափը չափազանց մեծ է։')
  book=openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
 except (zipfile.BadZipFile, KeyError, OSError) as e:
  raise ValueError('Հնարավոր չէ կարդալ Excel ֆայլը։') from e
 candidates=[]
 try:
  for sheet in book:
   if sheet.max_row>2000 or sheet.max_column>200: continue
   rows=list(sheet.iter_rows())
   totals=[]
   for row in rows[:6]:
    for cell in row:
     text=norm(cell.value)
     if text in ('обн. план','обновленный план'): totals.append((3,cell.column))
     elif text in ('всего','всего на 2026г.','всего на 2026г','итого 2026г.'): totals.append((2,cell.column))
   if not totals: continue
   col=max(totals)[1]; found={}; warnings=[]
   for row in rows:
    labels=[norm(c.value) for c in row[:4] if isinstance(c.value,str)]
    for key,(title,unit,aliases) in METRICS.items():
     if key in found or not any(a in labels for a in aliases): continue
     cell=row[col-1] if len(row)>=col else None
     if cell and isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool) and math.isfinite(cell.value):
      found[key]={'value':cell.value,'cell':cell.coordinate,'title':title,'unit':unit}
     else: warnings.append(title+'՝ տարեկան արժեքը բացակայում է։ Excel-ում վերահաշվարկեք և պահպանեք ֆայլը։')
   if 'cu' not in found or 'processed' not in found: continue
   status=rows[3] if len(rows)>3 else []
   actual_cols=[c.column for c in status if norm(c.value)=='факт']
   plan_cols=[c.column for c in status if norm(c.value) in ('план','опер. план','ожидание') and c.column<col and not (c.column<len(status) and norm(status[c.column].value)=='факт')]
   split={}
   # A split is valid only when the selected month values reproduce the annual total.
   for key in ('cu','mo','processed'):
    if key not in found or not actual_cols: continue
    rownum=sheet[found[key]['cell']].row
    vals=[sheet.cell(rownum,c).value for c in actual_cols+plan_cols]
    if all(isinstance(v,(int,float)) for v in vals) and math.isclose(sum(vals),found[key]['value'],rel_tol=1e-8,abs_tol=1e-5):
     split[key]={'actual':sum(vals[:len(actual_cols)]),'plan':sum(vals[len(actual_cols):]),'actual_months':len(actual_cols)}
   candidates.append({'filename':filename,'sheet':sheet.title,'metrics':found,'split':split,'warnings':warnings})
 finally: book.close()
 if not candidates: raise ValueError('Չգտնվեց աջակցվող պլանի թերթ՝ տարեկան գումարներով։ Աջակցվում են տրամադրված ԶՄՄԿ պլանների ձևաչափերը։')
 return max(candidates,key=lambda x:len(x['metrics']))

def compare(a,b):
 result=[]
 for key,(title,unit,_) in METRICS.items():
  if key not in a['metrics'] or key not in b['metrics']: continue
  x=a['metrics'][key]['value'];y=b['metrics'][key]['value']
  result.append({'key':key,'title':title,'unit':unit,'base':x,'updated':y,'delta':y-x,'percent':(y/x-1)*100 if x else None})
 factors=None
 keys=('cu','processed','cu_grade','cu_recovery')
 if all(k in p['metrics'] for p in (a,b) for k in keys):
  av={k:a['metrics'][k]['value'] for k in keys};bv={k:b['metrics'][k]['value'] for k in keys}
  def modeled(v):return v['processed']*v['cu_grade']/100*v['cu_recovery']/100
  if all(v[k]>0 for v in (av,bv) for k in keys):
   volume=av['cu']*(bv['processed']/av['processed']-1)
   grade=(av['cu']+volume)*(bv['cu_grade']/av['cu_grade']-1)
   recovery=bv['cu']-av['cu']-volume-grade
   factors=[{'title':'Վերամշակման ծավալ','value':volume},{'title':'Cu պարունակություն','value':grade},{'title':'Կորզում','value':recovery}]
 return {'base':a,'updated':b,'comparison':result,'factors':factors}

class Handler(BaseHTTPRequestHandler):
 def do_GET(self):
  if self.path not in ('/','/index.html'): self.send_error(404);return
  self.reply(200,(Path(__file__).parent/'index.html').read_bytes(),'text/html; charset=utf-8')
 def reply(self,status,data,kind='application/json; charset=utf-8'):
  self.send_response(status);self.send_header('Content-Type',kind);self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.end_headers();self.wfile.write(data)
 def do_POST(self):
  if self.path!='/api/compare':self.send_error(404);return
  try:
   length=int(self.headers.get('Content-Length','0'))
   if not 0<length<=57*1024*1024:raise ValueError('Ընտրեք մինչև 20 ՄԲ չափով երկու ֆայլ։')
   obj=json.loads(self.rfile.read(length))
   plans=[read_plan(base64.b64decode(obj[k]['data'],validate=True),obj[k]['name']) for k in ('base','updated')]
   self.reply(200,json.dumps(compare(*plans),ensure_ascii=False).encode())
  except Exception as e:
   message=str(e) if isinstance(e,ValueError) else 'Ֆայլի մշակումը չհաջողվեց։ Ստուգեք .xlsx ֆայլի կառուցվածքը։'
   self.reply(400,json.dumps({'error':message},ensure_ascii=False).encode())

if __name__=='__main__':
 host=os.environ.get('HOST','127.0.0.1');port=int(os.environ.get('PORT','8501'))
 print(f'Open http://{host}:{port}',flush=True)
 ThreadingHTTPServer((host,port),Handler).serve_forever()
