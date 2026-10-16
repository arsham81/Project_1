import io
import unittest
import openpyxl
from app import read_plan,compare

def fixture(updated=False,missing=False):
 w=openpyxl.Workbook();s=w.active
 s['A3']='Показатель';s['B3']='ед.изм.';s['C4']='Факт';s['D4']='План';s['E3']='Всего';s['F4']='Обн. План'
 entries=[('Итого переработка руды',100,120),('Среднее содержание Cu в руде',1,1.2),('Извлечение меди (Cu)',80,90),('Количество меди (Cu) в медном концентрате',.8,1.296)]
 for r,(label,x,y) in enumerate(entries,6):
  s.cell(r,1,label);s.cell(r,5,x);s.cell(r,6,y);s.cell(r,3,y*.75);s.cell(r,4,y*.25)
 if missing:s['F9']='=1+1'
 b=io.BytesIO();w.save(b);return read_plan(b.getvalue(),'test.xlsx')

class Tests(unittest.TestCase):
 def test_updated_total_preferred_and_split(self):
  p=fixture();self.assertEqual(p['metrics']['cu']['value'],1.296);self.assertEqual(p['metrics']['cu']['cell'],'F9');self.assertAlmostEqual(p['split']['cu']['actual'],.972)
 def test_factor_sum_and_zero_base(self):
  a=fixture();b=fixture();a['metrics']['cu']['value']=.8;a['metrics']['processed']['value']=100;a['metrics']['cu_grade']['value']=1;a['metrics']['cu_recovery']['value']=80
  r=compare(a,b);self.assertAlmostEqual(sum(x['value'] for x in r['factors']),.496)
  a['metrics']['cu']['value']=0;r=compare(a,b);self.assertIsNone(next(x for x in r['comparison'] if x['key']=='cu')['percent']);self.assertIsNone(r['factors'])
 def test_missing_formula_cache_rejected(self):
  with self.assertRaises(ValueError):fixture(missing=True)
 def test_invalid_file(self):
  with self.assertRaises(ValueError):read_plan(b'not a workbook','bad.xlsx')
 def test_inconsistent_split_omitted(self):
  p=fixture();self.assertNotIn('cu_grade',p['split'])

if __name__=='__main__':unittest.main()
